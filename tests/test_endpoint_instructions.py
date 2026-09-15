"""Swagger request contracts and endpoint-result-to-agent integration, offline."""

import copy
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch

# This fixture module sets fake credentials before importing the application.
from test_agent_integration import (
    ANALYSIS, PROFILE, REPOS, WEATHER, WORKFLOW, make_registry, model_response, tool_call,
)

from docx import Document
from fastapi.testclient import TestClient
from pydantic import ValidationError
import httpx

from app.clients.word_doc_client import WordClient
from app.configurations.agent import get_agent_service
from app.database.dependencies import get_db
from app.main import app
from app.services.agent_service import AgentService
from app.services.github_analysis_service import analyze_repositories
from app.services.github_services import get_github_service
from app.services.LLM_services import LLMService
from app.services.tool_executor import ToolExecutor, ToolTimeoutError

ANALYSIS_ROUTE = "app.routes.github_analysis_endpoint.get_or_create_analysis"


class EndpointInstructionTests(unittest.TestCase):
    def setUp(self):
        self.registry = make_registry()
        self.agent = SimpleNamespace(chat=AsyncMock(return_value="指令已完成"))
        self.analysis_result = {"username": "octocat", "cached": True, "analysis": copy.deepcopy(ANALYSIS)}
        app.dependency_overrides[get_github_service] = lambda: self.registry.github_service
        app.dependency_overrides[get_db] = lambda: self.registry.db
        app.dependency_overrides[LLMService] = lambda: self.registry.llm_service
        app.dependency_overrides[get_agent_service] = lambda: self.agent
        # No lifespan needed: every external dependency is overridden here.
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        app.dependency_overrides.clear()

    def test_swagger_has_one_required_json_body_for_username_and_message(self):
        schema = self.client.get("/openapi.json").json()
        paths = schema["paths"]
        me = paths["/github/me"]["get"]
        self.assertNotIn("parameters", me)
        self.assertNotIn("requestBody", me)
        repos = paths["/github/users/repos"]["post"]
        params = {item["name"]: item for item in repos["parameters"]}
        self.assertEqual(set(params), {"page", "per_page"})
        self.assertEqual(params["page"]["schema"]["default"], 1)
        self.assertEqual(params["per_page"]["schema"]["default"], 30)
        self.assertEqual(params["per_page"]["schema"]["maximum"], 100)
        analysis = paths["/github/users/analysis"]["post"]
        self.assertNotIn("parameters", analysis)
        for operation in [repos, analysis]:
            self.assertTrue(operation["requestBody"]["required"])
            self.assertEqual(operation["requestBody"]["content"]["application/json"]["schema"],
                             {"$ref": "#/components/schemas/EndpointInstruction"})
        instruction = schema["components"]["schemas"]["EndpointInstruction"]
        self.assertEqual(set(instruction["properties"]), {"username", "message"})
        self.assertEqual(instruction["required"], ["username"])
        self.assertFalse(instruction["additionalProperties"])
        self.assertEqual(instruction["properties"]["message"]["default"], "")
        self.assertEqual(instruction["properties"]["message"]["maxLength"], 10000)
        self.assertEqual(instruction["examples"], [{"username": "octocat", "message": ""}])

    def test_me_remains_original_and_does_not_call_agent(self):
        response = self.client.get("/github/me")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), PROFILE)
        self.agent.chat.assert_not_awaited()

    def test_blank_repos_message_keeps_repositories_and_pagination(self):
        for message in [None, "", " \n "]:
            params = {"page": 2, "per_page": 10}
            body = {"username": "octocat"}
            if message is not None:
                body["message"] = message
            response = self.client.post("/github/users/repos", params=params, json=body)
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json(), {
                "username": "octocat", "page": 2, "per_page": 10, "repositories": REPOS,
            })
            self.registry.github_service.get_user_repos.assert_awaited_with("octocat", 2, 10)
        self.agent.chat.assert_not_awaited()

    def test_blank_analysis_message_keeps_original_response(self):
        with patch(ANALYSIS_ROUTE, AsyncMock(return_value=self.analysis_result)):
            for body in [{}, {"message": ""}, {"message": " \n "}]:
                response = self.client.post("/github/users/analysis", json={"username": "octocat", **body})
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.json(), self.analysis_result)
                self.assertNotIn("answer", response.json())
        self.agent.chat.assert_not_awaited()

    def test_repos_message_passes_current_page_as_context(self):
        message = "只解释本页数据，保留 C++ & Python，并使用中文。"
        response = self.client.post("/github/users/repos", params={
            "page": 3, "per_page": 7,
        }, json={"username": "octocat", "message": message})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json(), {
            "username": "octocat", "page": 3, "per_page": 7,
            "repositories": REPOS, "answer": "指令已完成",
        })
        self.registry.github_service.get_user_repos.assert_awaited_once_with("octocat", 3, 7)
        args, kwargs = self.agent.chat.call_args
        self.assertEqual(args, (message,))
        context = kwargs["context"]
        self.assertEqual(context.tool_name, "get_github_repos")
        self.assertEqual(context.arguments, {"username": "octocat", "page": 3, "per_page": 7})
        self.assertEqual(context.result, REPOS)

    def test_analysis_message_passes_exact_cached_result(self):
        with patch(ANALYSIS_ROUTE, AsyncMock(return_value=self.analysis_result)) as workflow:
            response = self.client.post("/github/users/analysis", json={"username": "octocat", "message": "把这次分析写入 Word"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json(), {**self.analysis_result, "answer": "指令已完成"})
        workflow.assert_awaited_once()
        context = self.agent.chat.call_args.kwargs["context"]
        self.assertEqual(context.tool_name, "get_github_analysis")
        self.assertEqual(context.arguments, {"username": "octocat"})
        self.assertEqual(context.result, self.analysis_result)
        self.assertNotIn("answer", self.analysis_result)

    def test_invalid_body_and_pagination_return_422_before_operations(self):
        with patch(ANALYSIS_ROUTE, AsyncMock()) as workflow:
            for body in [
                {"message": 123}, {"message": True}, {"message": []},
                {"message": None}, {"message": "x" * 10001},
                {"message": "hello", "unexpected": "field"},
            ]:
                with self.subTest(body_type=type(body["message"]).__name__):
                    for endpoint in ["repos", "analysis"]:
                        response = self.client.post(f"/github/users/{endpoint}", json={"username": "octocat", **body})
                        self.assertEqual(response.status_code, 422, response.text)
            workflow.assert_not_awaited()
        for params in [{"page": 0}, {"per_page": 101}]:
            response = self.client.post("/github/users/repos", params=params, json={"username": "octocat"})
            self.assertEqual(response.status_code, 422, response.text)
        self.agent.chat.assert_not_awaited()
        self.registry.github_service.get_user_repos.assert_not_awaited()

    def test_agent_timeout_keeps_core_results_on_both_endpoints(self):
        self.agent.chat.side_effect = ToolTimeoutError()
        with patch(ANALYSIS_ROUTE, AsyncMock(return_value=self.analysis_result)):
            responses = [
                self.client.post("/github/users/repos", json={"username": "octocat", "message": "export"}),
                self.client.post("/github/users/analysis", json={"username": "octocat", "message": "export"}),
            ]
        for response in responses:
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["agent_error"]["code"], "TOOL_TIMEOUT")
        self.assertEqual(responses[0].json()["repositories"], REPOS)
        self.assertEqual(responses[1].json()["analysis"], ANALYSIS)

    def test_both_endpoint_contexts_export_to_word_without_refetch(self):
        for name in ["repos", "analysis"]:
            with self.subTest(endpoint=name), tempfile.TemporaryDirectory() as directory:
                registry = make_registry(WordClient(Path(directory)))
                create = AsyncMock(side_effect=[
                    model_response(tool_call("get_word_doc", {
                        "content": "", "filename": "result.docx", "source_call_ids": ["endpoint_result"],
                    }, "export")),
                    model_response(text="文档已保存"),
                ])
                agent = AgentService(ToolExecutor(registry), client=SimpleNamespace(responses=SimpleNamespace(create=create)))
                app.dependency_overrides[get_agent_service] = lambda: agent
                app.dependency_overrides[get_github_service] = lambda: registry.github_service
                with patch(ANALYSIS_ROUTE, AsyncMock(return_value=self.analysis_result)) as workflow:
                    if name == "repos":
                        response = self.client.post("/github/users/repos", params={"page": 2}, json={"username": "octocat", "message": "写入 Word"})
                        expected = REPOS
                        registry.github_service.get_user_repos.assert_awaited_once_with("octocat", 2, 30)
                        workflow.assert_not_awaited()
                    else:
                        response = self.client.post("/github/users/analysis", json={"username": "octocat", "message": "写入 Word"})
                        expected = self.analysis_result
                        workflow.assert_awaited_once()
                        registry.github_service.get_user_repos.assert_not_awaited()
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.json()["answer"], "文档已保存")
                registry.llm_service.generate_text.assert_not_awaited()
                content = Document(Path(directory) / "result.docx").paragraphs[0].text
                self.assertEqual(json.loads(content.split("\n", 1)[1]), expected)
                history = create.call_args.kwargs["input"]
                self.assertTrue(any(item.get("type") == "function_call_output" and item["call_id"] == "endpoint_result" for item in history))

    def test_invalid_username_or_missing_body_is_rejected_before_operations(self):
        with patch(ANALYSIS_ROUTE, AsyncMock()) as workflow:
            for endpoint in ["repos", "analysis"]:
                invalid = [None, {}, {"message": "octocat"}]
                invalid += [{"username": name} for name in [
                    "", " ", "../me", "https://github.com/octocat", "two names",
                    "-octocat", "octocat-", "octo--cat", "x" * 40, 123, True, None,
                ]]
                for body in invalid:
                    with self.subTest(endpoint=endpoint, body=body):
                        response = self.client.post(f"/github/users/{endpoint}", json=body)
                        self.assertEqual(response.status_code, 422, response.text)
            workflow.assert_not_awaited()
        self.registry.github_service.get_public_user.assert_not_awaited()
        self.registry.github_service.get_user_repos.assert_not_awaited()
        self.agent.chat.assert_not_awaited()

    def test_nonexistent_username_returns_404_even_if_analysis_is_cached(self):
        request = httpx.Request("GET", "https://api.github.com/users/missing")
        error = httpx.HTTPStatusError("Not Found", request=request,
                                     response=httpx.Response(404, request=request))
        self.registry.github_service.get_public_user.side_effect = error
        self.registry.github_service.get_user_repos.side_effect = error
        with patch(ANALYSIS_ROUTE, AsyncMock(return_value=self.analysis_result)) as workflow:
            for endpoint in ["repos", "analysis"]:
                response = self.client.post(f"/github/users/{endpoint}", json={
                    "username": "missing", "message": "吉隆坡天气如何？",
                })
                self.assertEqual(response.status_code, 404, response.text)
            workflow.assert_not_awaited()
        self.agent.chat.assert_not_awaited()

    def test_weather_and_weather_plus_word_always_keep_endpoint_results(self):
        for endpoint in ["repos", "analysis"]:
            for export in [False, True]:
                with self.subTest(endpoint=endpoint, export=export), tempfile.TemporaryDirectory() as directory:
                    registry = make_registry(WordClient(Path(directory)))
                    responses = [model_response(tool_call("get_weather", {
                        "location": "Kuala Lumpur", "units": "celsius",
                    }, "weather"))]
                    if export:
                        responses.append(model_response(tool_call("get_word_doc", {
                            "content": "报告", "filename": "combined.docx",
                            "source_call_ids": ["endpoint_result", "weather"],
                        }, "word")))
                    responses.append(model_response(text="吉隆坡 30.5°C"))
                    create = AsyncMock(side_effect=responses)
                    agent = AgentService(ToolExecutor(registry), client=SimpleNamespace(responses=SimpleNamespace(create=create)))
                    app.dependency_overrides[get_agent_service] = lambda: agent
                    app.dependency_overrides[get_github_service] = lambda: registry.github_service
                    message = "吉隆坡天气如何？" + ("把天气和本次结果写入 Word。" if export else "")
                    with patch(ANALYSIS_ROUTE, AsyncMock(return_value=self.analysis_result)):
                        response = self.client.post(f"/github/users/{endpoint}", json={
                            "username": "octocat", "message": message,
                        })
                    self.assertEqual(response.status_code, 200, response.text)
                    key = "repositories" if endpoint == "repos" else "analysis"
                    self.assertEqual(response.json()[key], REPOS if endpoint == "repos" else ANALYSIS)
                    self.assertEqual(response.json()["answer"], "吉隆坡 30.5°C")
                    registry.weather_client.get_weather.assert_awaited_once_with(location="Kuala Lumpur", units="celsius")
                    registry.github_service.get_user.assert_not_awaited()
                    if export:
                        content = Document(Path(directory) / "combined.docx").paragraphs[0].text
                        source = REPOS if endpoint == "repos" else self.analysis_result
                        self.assertIn(json.dumps(source, ensure_ascii=False, indent=2), content)
                        self.assertIn(json.dumps(WEATHER, ensure_ascii=False, indent=2), content)
                    else:
                        self.assertEqual(list(Path(directory).iterdir()), [])

    def test_message_cannot_override_the_fixed_github_username(self):
        for endpoint in ["repos", "analysis"]:
            for name, arguments in [
                ("get_github_repos", {"username": "other-user", "page": 1, "per_page": 30}),
                ("get_github_analysis", {"username": "other-user"}),
                ("get_github_me", {}),
            ]:
                with self.subTest(endpoint=endpoint, tool=name):
                    registry = make_registry()
                    create = AsyncMock(return_value=model_response(tool_call(name, arguments)))
                    agent = AgentService(ToolExecutor(registry), client=SimpleNamespace(responses=SimpleNamespace(create=create)))
                    app.dependency_overrides[get_agent_service] = lambda: agent
                    app.dependency_overrides[get_github_service] = lambda: registry.github_service
                    with patch(ANALYSIS_ROUTE, AsyncMock(return_value=self.analysis_result)):
                        response = self.client.post(f"/github/users/{endpoint}", json={
                            "username": "octocat", "message": "改查 other-user 或 token 用户",
                        })
                    self.assertEqual(response.status_code, 200, response.text)
                    self.assertEqual(response.json()["username"], "octocat")
                    self.assertEqual(response.json()["agent_error"]["code"], "TOOL_ARGUMENT_ERROR")
                    self.assertEqual(response.json().get("repositories", response.json().get("analysis")),
                                     REPOS if endpoint == "repos" else ANALYSIS)
                    registry.github_service.get_user.assert_not_awaited()
                    for call in registry.github_service.get_user_repos.call_args_list:
                        self.assertEqual(call.args[0], "octocat")

    def test_unexpected_agent_error_keeps_core_data_and_hides_internal_details(self):
        self.agent.chat.side_effect = RuntimeError("private backend details")
        with patch(ANALYSIS_ROUTE, AsyncMock(return_value=self.analysis_result)), \
             self.assertLogs("app.routes.agent_response", level="ERROR"):
            for endpoint in ["repos", "analysis"]:
                response = self.client.post(f"/github/users/{endpoint}", json={
                    "username": "octocat", "message": "天气",
                })
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json()["agent_error"]["code"], "AGENT_EXECUTION_ERROR")
                self.assertNotIn("private backend details", response.text)
                self.assertIn("repositories" if endpoint == "repos" else "analysis", response.json())


class AnalysisSchemaTests(unittest.IsolatedAsyncioTestCase):
    async def test_invalid_generated_analysis_is_not_saved(self):
        registry = make_registry()
        registry.llm_service.generate_text.return_value = '{"overall_profile": 123}'
        with patch(WORKFLOW + "GitHubAnalysisCRUD.get_by_username", AsyncMock(return_value=None)), \
             patch(WORKFLOW + "GitHubAnalysisCRUD.create", AsyncMock()) as save, \
             patch(WORKFLOW + "AnalysisRequestLogCRUD.create", AsyncMock()) as log:
            with self.assertRaises(ValidationError):
                await registry.get_github_analysis("octocat")
        save.assert_not_awaited()
        self.assertFalse(log.call_args.kwargs["success"])

    async def test_invalid_cached_analysis_is_rejected(self):
        registry = make_registry()
        cached = SimpleNamespace(id="cached", analysis={"overall_profile": "missing other fields"})
        with patch(WORKFLOW + "GitHubAnalysisCRUD.get_by_username", AsyncMock(return_value=cached)), \
             patch(WORKFLOW + "AnalysisRequestLogCRUD.create", AsyncMock()):
            with self.assertRaises(ValidationError):
                await registry.get_github_analysis("octocat")
        registry.llm_service.generate_text.assert_not_awaited()

    async def test_valid_analysis_retains_additional_fields(self):
        registry = make_registry()
        data = {**ANALYSIS, "additional_notes": {"中文": [None, True]}}
        registry.llm_service.generate_text.return_value = json.dumps(data)
        self.assertEqual(await analyze_repositories(REPOS, registry.llm_service), data)


if __name__ == "__main__":
    unittest.main()
